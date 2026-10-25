"""Draw one shared annotation set and write an Excel sheet per annotator (docs/ANOTACJA.md).
    python scripts/annotation_sample.py --name pilot_v1 --n-per-platform 100 --annotators A B C D
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.lexicon.slurs import TERMS, find_slurs  # noqa: E402
from src.preprocessing.text import QUOTE_RE, normalize_whitespace  # noqa: E402

OUT_ROOT = ROOT / 'data/annotation'

LABELS = {'hate': ['1', '0'], 'hate_migrants': ['1', '0'], 'uncertain': ['1']}
SHEET_COLUMNS = ['item_id', 'context_post', 'context_parent', 'context_quote', 'text', 'hate', 'hate_migrants',
                 'uncertain', 'note']
POOLS = {'mig_hostile': 0.4, 'mig_other': 0.4, 'ambiguous': 0.2}

_PL = str.maketrans('ąćęłńóśźżĄĆĘŁŃÓŚŹŻ', 'acelnoszzACELNOSZZ')
# migrants and Ukrainians as people; 'Ukraina' alone (the state, the war) does not count
MIGRATION = re.compile(
    r'\b(?:i?migra\w*|imigr\w*|emigrant\w*|uchodzc\w*|uchodzstw\w*|nachodzc\w*|azyl\w*|cudzoziem\w*|obcokraj\w*|'
    r'relokac\w*|deport\w*|ukrai?nc\w*|ukrainiec|ukraink\w*|ukrop\w*|ukr(?:y|ow|om|ami)\b|banderow\w*|'
    r'muzulman\w*|muzul\w*|islam\w*|afryk\w*|arab\w*|murzyn\w*|ciapa\w*|syryj\w*|afgan\w*|hindus\w*|gruzin\w*|'
    r'kolumbij\w*|ubogac\w*|inzynier(?!i)\w*|multikulti|kebab\w*|sniad\w*)')
HOSTILE = re.compile(
    r'\b(?:wypierdala\w*|wypierdol\w*|spierdala\w*|won\b|precz|wynocha|wywal\w*|wyrzuc\w*|odesla\w*|'
    r'kurw\w*|jeba\w*|jebi\w*|pierdol\w*|chuj\w*|huj\w*|debil\w*|idiot\w*|kretyn\w*|bydl\w*|scierw\w*|smiec\w*|'
    r'robact\w*|szarancz\w*|pasozyt\w*|darmozjad\w*|dzicz\w*|dzikus\w*|malp\w*|zwierz\w*|ciapa\w*|brudas\w*|'
    r'czarnuch\w*|nachodzc\w*|ukrop\w*|ukr(?:y|ow|om|ami)\b|banderow\w*|ubogac\w*|inzynier(?!i)\w*|faszys\w*|'
    r'rasis\w*|onuc\w*|kacap\w*|lewac\w*|lewak\w*|zabic\w*|strzela\w*|spali\w*|wieszac\w*|gwalc\w*|zlodzie\w*)')

INSTRUCTION = [
    'Zasady: docs/CODEBOOK.md (wersja {version}). Anotator: {annotator}.',
    '',
    'hate = 1          autor atakuje kogokolwiek: obelga, dehumanizacja, stereotyp, wykluczenie, przemoc, kod/ironia.',
    'hate_migrants = 1  ten atak jest skierowany do migrantów (w tym Ukraińców) ze względu na pochodzenie lub status.',
    'Dozwolone pary hate/hate_migrants: 0/0, 1/0, 1/1.',
    '',
    'Kolejność decyzji:',
    '1. Czy AUTOR (nie cytat) atakuje jakąś osobę lub grupę?        nie -> 0/0, tak -> hate = 1',
    '2. Czy atak dotyczy migrantów ze względu na pochodzenie/status,',
    '   a nie czyn konkretnej osoby?                                tak -> hate_migrants = 1, nie -> 0',
    '',
    'Komentarze z obelgą etniczną mają już wpisane hate = 1 i hate_migrants = 1',
    'Zmień je tylko przy fałszywym trafieniu (inne znaczenie słowa, nazwa własna)',
    'i wtedy opisz to w note.',
    '',
    'Kontekst (context_post, context_parent, context_quote) służy tylko do ustalenia adresata.',
    'context_quote to cytat z innego komentarza - nie jest słowami autora i nie wpływa na kody.',
    'Jeśli masz wątpliwości: wybierz najbardziej oczywiste odczytanie, wpisz uncertain = 1 i krótko opisz w note.',
    'Nie zmieniaj kolumn item_id, context_*, text ani nazwy pliku.',
]


def fold(text):
    return str(text).translate(_PL).lower()


def split_quote(text):
    """Author's own words and the quoted lines ('>') as two strings."""
    quote = '\n'.join(m.group(0)[len(m.group(1)):].lstrip('> \t') for m in QUOTE_RE.finditer(text))
    return normalize_whitespace(QUOTE_RE.sub(r'\1', text)), quote


def read_csv(path):
    return pd.read_csv(path, encoding='utf-8-sig', dtype=str, keep_default_na=False)


def load_platform(platform):
    """Comments with post and parent context"""
    rows = []
    for raw_dir in sorted((ROOT / 'data/raw' / platform).glob('*/')):
        clean = ROOT / 'data/interim' / platform / raw_dir.name / 'clean'
        src = clean if (clean / 'comments.csv').exists() else raw_dir
        if not (src / 'comments.csv').exists():
            continue
        c = read_csv(src / 'comments.csv')
        posts = read_csv(raw_dir / 'posts.csv') if (raw_dir / 'posts.csv').exists() else pd.DataFrame()
        raw = c['text_raw'] if 'text_raw' in c else c['text']    # masked original, with emoji
        c['shown'], c['context_quote'] = zip(*raw.map(split_quote)) if len(c) else ([], [])
        by_id = dict(zip(c['comment_id'], c['shown']))
        c['context_post'] = c['post_id'].map(dict(zip(posts.get('post_id', []), posts.get('text', [])))).fillna('')
        c['context_parent'] = c['reply_to'].map(by_id).fillna('') if 'reply_to' in c else ''
        c['post_migration'] = c['context_post'].map(lambda t: bool(MIGRATION.search(fold(t))))
        c['source'] = 'clean' if src == clean else 'raw'
        flags = clean.parent / 'slurs' / 'comments.csv'   # stage slurs_<platform>, same rows as clean
        if src == clean and flags.exists():
            s = read_csv(flags)
            c['slurs'] = c['comment_id'].map(dict(zip(s['comment_id'], s['slurs'].map(mark_context))))
        rows.append(c.assign(platform=platform, event_folder=raw_dir.name))
    return pd.concat(rows, ignore_index=True) if rows else None


def load_youtube_csv(path):
    """YouTube export from 02_youtube_EDA (event_name, video_title, text, optional comment id)."""
    y = read_csv(path)
    missing = {'event_name', 'video_title', 'text'} - set(y.columns)
    if missing:
        sys.exit(f'{path}: missing columns {sorted(missing)}')
    id_col = next((c for c in ('comment_id', 'id') if c in y.columns), None)
    y['comment_id'] = y[id_col] if id_col else 'yt_row_' + y.index.astype(str)
    y['post_id'] = y.get('video_id', '')
    y['shown'], y['context_quote'] = zip(*y['text'].map(split_quote)) if len(y) else ([], [])
    y['context_post'], y['context_parent'] = y['video_title'], ''
    y['post_migration'] = y['context_post'].map(lambda t: bool(MIGRATION.search(fold(t))))
    return y.assign(platform='youtube', event_folder=y['event_name'], source='csv')


_CONTEXT = {t.name: t.needs_context for t in TERMS}


def mark_context(names):
    """'ukry; bandera' -> 'ukry; bandera (?)'; (?) marks terms that depend on context."""
    return '; '.join(n + (' (?)' if _CONTEXT.get(n) else '') for n in names.split('; ') if n)


def slur_label(text):
    return mark_context('; '.join(t.name for t in find_slurs(text)))


def pick(df, n, seed):
    """40% migration + hostile word, 40% migration only, 20% hostile word under a migration post.
    Comments with a slur stay in the draw and count as migration + hostile."""
    df = df[df['shown'].str.len().between(20, 1200)]
    df = df[df['shown'].map(lambda t: len(re.findall(r'[^\W\d_]{2,}', t)) >= 3)].drop_duplicates('shown')
    if 'slurs' not in df:
        df = df.assign(slurs=pd.NA)
    todo = df['slurs'].isna()   # events without slurs/ output (stage not run yet)
    df = df.assign(slurs=df['slurs'].where(~todo, df.loc[todo, 'shown'].map(slur_label)))
    f = df['shown'].map(fold)
    slur = df['slurs'] != ''
    mig = f.map(lambda t: bool(MIGRATION.search(t))) | slur
    hostile = f.map(lambda t: bool(HOSTILE.search(t))) | slur
    pools = {'mig_hostile': df[mig & hostile], 'mig_other': df[mig & ~hostile],
             'ambiguous': df[~mig & hostile & df['post_migration']]}
    want = {k: round(n * share) for k, share in POOLS.items()}
    want['mig_other'] += n - sum(want.values())
    want['mig_other'] += sum(max(0, want[k] - len(pools[k])) for k in pools if k != 'mig_other')
    return pd.concat(pools[k].sample(n=min(want[k], len(pools[k])), random_state=seed).assign(pool=k) for k in pools)


def prefilled(slurs, mode):
    """'1' for hate and hate_migrants when the sheet should start at 1/1 (codebook pkt 4.12)."""
    if not slurs or mode == 'none':
        return ''
    return '1' if mode == 'all' or any(not s.endswith('(?)') for s in slurs.split('; ')) else ''


def write_sheet(items, path, annotator, version, prefill='strict'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'anotacja'
    ws.append(SHEET_COLUMNS)
    for _, r in items.iterrows():
        auto = prefilled(r['slurs'], prefill)
        ws.append([r['item_id'], r['context_post'][:600], r['context_parent'][:600], r['context_quote'][:600],
                   r['shown'], auto, auto, '', ''])
    n = len(items) + 1
    widths = {'item_id': 9, 'context_post': 40, 'context_parent': 35, 'context_quote': 35, 'text': 70, 'hate': 8,
              'hate_migrants': 14, 'uncertain': 11, 'note': 35}
    for i, col in enumerate(SHEET_COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        ws.column_dimensions[cell.column_letter].width = widths[col]
        cell.font = Font(bold=True)
        if col in LABELS:
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
            dv = DataValidation(type='list', formula1='"' + ','.join(LABELS[col]) + '"', allow_blank=True,
                                showErrorMessage=True, error='Wybierz wartość z listy (docs/CODEBOOK.md).')
            ws.add_data_validation(dv)
            dv.add(f'{cell.column_letter}2:{cell.column_letter}{n}')
        for row in range(2, n + 1):
            ws.cell(row=row, column=i).alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    info = wb.create_sheet('instrukcja')
    for line in INSTRUCTION:
        info.append([line.format(version=version, annotator=annotator)])
    info.column_dimensions['A'].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def codebook_version():
    cb = ROOT / 'docs/CODEBOOK.md'
    m = re.search(r'Wersja ([\d.]+)', cb.read_text(encoding='utf-8')) if cb.exists() else None
    return m.group(1) if m else '?'


def main():
    ap = argparse.ArgumentParser(description='Draw a shared annotation set.')
    ap.add_argument('--name', required=True, help='annotation round, e.g. pilot_v1')
    ap.add_argument('--n-per-platform', type=int, default=100)
    ap.add_argument('--annotators', nargs='+', default=['A', 'B', 'C', 'D'])
    ap.add_argument('--platforms', nargs='+', default=['wykop', 'bluesky', 'youtube'])
    ap.add_argument('--youtube-csv', help='YouTube CSV if data/raw/youtube does not exist yet')
    ap.add_argument('--seed', type=int, default=2026)
    ap.add_argument('--prefill', choices=['strict', 'all', 'none'], default='strict',
                    help='1/1 for slur matches: all, strict (skip context-only terms) or none')
    ap.add_argument('--append', action='store_true',
                    help='add items (e.g. a new platform) to an existing round; old item_ids stay')
    ap.add_argument('--force', action='store_true')
    args = ap.parse_args()

    out = OUT_ROOT / args.name
    existing = None
    if (out / 'manifest.csv').exists():
        if args.append:
            existing = read_csv(out / 'manifest.csv').rename(columns={'text': 'shown'})
        elif not args.force:
            sys.exit(f'{out} already exists - annotations refer to its item_ids. '
                     'Use another --name, or --append to add items to it.')

    picked = []
    for platform in args.platforms:
        df = load_youtube_csv(args.youtube_csv) if platform == 'youtube' and args.youtube_csv else load_platform(platform)
        if df is None or df.empty:
            print(f'WARNING: no data for {platform} (run `dvc pull`'
                  f'{" or pass --youtube-csv" if platform == "youtube" else ""}); skipped')
            continue
        s = pick(df, args.n_per_platform, args.seed)
        print(f'{platform}: {len(df):,} comments -> {len(s)} picked {s["pool"].value_counts().to_dict()}, '
              f'{(s["slurs"] != "").sum()} with a slur')
        picked.append(s)
    if not picked:
        sys.exit('Nothing to sample.')

    items = pd.concat(picked).sample(frac=1, random_state=args.seed).reset_index(drop=True)
    start = 0
    if existing is not None:
        seen = set(zip(existing['platform'], existing['comment_id']))
        items = items[[k not in seen for k in zip(items['platform'], items['comment_id'])]]
        start = int(existing['item_id'].str[1:].astype(int).max())
        print(f'appending {len(items)} items to the {len(existing)} already in {args.name}')
    items['item_id'] = [f'P{start + i + 1:04d}' for i in range(len(items))]
    if existing is not None:
        items = pd.concat([existing, items], ignore_index=True).fillna('')
    out.mkdir(parents=True, exist_ok=True)
    items[['item_id', 'platform', 'event_folder', 'comment_id', 'post_id', 'pool', 'source', 'context_post',
           'context_parent', 'context_quote', 'shown', 'slurs']].rename(columns={'shown': 'text'}) \
        .to_csv(out / 'manifest.csv', index=False, encoding='utf-8-sig')
    version = codebook_version()
    for a in args.annotators:
        write_sheet(items, out / 'sheets' / f'{args.name}_{a}.xlsx', a, version, args.prefill)
    (out / 'annotated').mkdir(exist_ok=True)
    print(f'\n{len(items)} comments for {len(args.annotators)} annotators -> {out / "sheets"}')
    print(f'Filled sheets go to {out / "annotated"}, then: python scripts/annotation_agreement.py --name {args.name}')


if __name__ == '__main__':
    main()
